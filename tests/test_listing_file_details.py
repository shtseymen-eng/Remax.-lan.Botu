import csv

from openpyxl import load_workbook

from remax_bot.core import Listing
from remax_bot.importer import read_list_file, write_ai_template, write_list_file


def detailed_listing():
    return Listing(
        "1234567890",
        "Yahya Kaptan Site İçinde Satılık Daire",
        "https://www.sahibinden.com/ilan/x-1234567890/detay",
        "Ayşe Yılmaz",
        "0532 111 22 33",
        "5.750.000 TL",
        "Kocaeli / İzmit / Yahya Kaptan",
        "3+1",
        "Satılık",
        "Daire",
        "145 m²",
        "26 Ağustos 2026",
        "https://carsigayrimenkulkocaeli.sahibinden.com/",
        "Kocaeli",
        "İzmit",
        "Yahya Kaptan",
        "Akasyalar Caddesi",
        "Zeray Sitesi",
        "Yahya Kaptan Mah. Akasyalar Cad. No: 10",
        "Ara kat, kapalı otoparklı satılık daire.",
        "130 m²",
        {"Isıtma": "Doğalgaz Kombi", "Bulunduğu Kat": "3", "Otopark": "Kapalı"},
    )


def assert_detailed_listing(item):
    expected = detailed_listing()
    assert item == expected


def test_xlsx_round_trip_preserves_fixed_and_dynamic_listing_details(tmp_path):
    path = tmp_path / "ilanlar.xlsx"

    write_list_file(path, [detailed_listing()])

    [item] = read_list_file(path)
    assert_detailed_listing(item)
    workbook = load_workbook(path, data_only=True)
    sheet = workbook["İlan Listesi"]
    headers = [cell.value for cell in sheet[1]]
    link_column = headers.index("İlan Linki") + 1
    assert sheet.cell(2, link_column).hyperlink.target.endswith("/detay")
    assert headers[-3:] == ["Bulunduğu Kat", "Isıtma", "Otopark"]


def test_csv_round_trip_preserves_turkish_text_and_dynamic_details(tmp_path):
    path = tmp_path / "ilanlar.csv"

    write_list_file(path, [detailed_listing()])

    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    [item] = read_list_file(path)
    assert_detailed_listing(item)


def test_unknown_nonempty_columns_are_imported_as_details_and_blank_columns_are_ignored(tmp_path):
    path = tmp_path / "ai.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=[
                "İlan Linki",
                "Başlık",
                "İl",
                "İlçe",
                "Mahalle",
                "Isıtma",
                "Bulunduğu Kat",
                "Boş Özellik",
            ],
            delimiter=";",
        )
        writer.writeheader()
        writer.writerow(
            {
                "İlan Linki": "https://www.sahibinden.com/ilan/x-1234567890/detay",
                "Başlık": "Türkçe Karakterli İlan",
                "İl": "Kocaeli",
                "İlçe": "İzmit",
                "Mahalle": "Yahya Kaptan",
                "Isıtma": "Yerden Isıtma",
                "Bulunduğu Kat": "Ara Kat",
                "Boş Özellik": "",
            }
        )

    [item] = read_list_file(path)

    assert item.province == "Kocaeli"
    assert item.district == "İzmit"
    assert item.neighborhood == "Yahya Kaptan"
    assert item.details == {"Isıtma": "Yerden Isıtma", "Bulunduğu Kat": "Ara Kat"}


def test_old_bolge_column_still_populates_legacy_location(tmp_path):
    path = tmp_path / "old.csv"
    path.write_text(
        "İlan Linki;Başlık;Bölge\nhttps://example.test/ilan/123456;Eski İlan;Kocaeli / İzmit\n",
        encoding="utf-8-sig",
    )

    [item] = read_list_file(path)

    assert item.location == "Kocaeli / İzmit"


def test_ai_template_has_input_headers_and_plain_language_instructions(tmp_path):
    path = tmp_path / "REMAX İlan Bot Şablon.xlsx"

    write_ai_template(path)

    workbook = load_workbook(path, data_only=True)
    assert workbook.sheetnames == ["İlanlar", "Alan Açıklamaları"]
    headers = [cell.value for cell in workbook["İlanlar"][1]]
    assert headers[:4] == ["İlan No", "Başlık", "Danışman", "Telefon"]
    assert "İlan Linki" in headers
    assert "Açık Adres" in headers
    assert "Isıtma" in headers
    instruction_text = " ".join(
        str(cell.value or "") for row in workbook["Alan Açıklamaları"] for cell in row
    )
    assert "yalnızca İlan Linki zorunludur" in instruction_text
    assert "tahmin etmeyin" in instruction_text
